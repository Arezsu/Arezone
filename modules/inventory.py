from __future__ import annotations

import sqlite3
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from database.db import now_text
from modules.formatters import money, number, parse_amount
from modules.product_codes import split_primary_and_alternates
from modules.product_search import search_products
from modules.security import get_setting, verify_secret
from modules.theme import fit_window
from modules.ui_fx import ask_secret, pop_in_window


class InventoryWindow(tk.Toplevel):
    def __init__(self, master, conn: sqlite3.Connection, on_import: Callable[[], None] | None = None) -> None:
        super().__init__(master)
        self.conn = conn
        self.on_import = on_import
        self.query_var = tk.StringVar()

        self.title("Inventario AREZONE")
        fit_window(self, 1080, 680, min_width=960, min_height=600)
        self.transient(master)
        self.grab_set()
        self.columnconfigure(0, weight=1)
        self.rowconfigure(1, weight=1)

        self._build()
        self.search()

    def _build(self) -> None:
        top = ttk.Frame(self, padding=10)
        top.grid(row=0, column=0, sticky="ew")
        top.columnconfigure(0, weight=1)
        entry = ttk.Entry(top, textvariable=self.query_var, font=("Segoe UI", 12))
        entry.grid(row=0, column=0, sticky="ew", padx=(0, 8), ipady=5)
        entry.bind("<Return>", lambda _event: self.search())
        ttk.Button(top, text="Buscar", command=self.search).grid(row=0, column=1, padx=3)
        ttk.Button(top, text="Migrar Excel", command=self.import_inventory).grid(row=0, column=2, padx=3)
        ttk.Button(top, text="Exportar Excel", command=self.export_inventory).grid(row=0, column=3, padx=3)

        columns = ("code", "name", "brand", "category", "stock", "cost", "public", "tax")
        self.tree = ttk.Treeview(self, columns=columns, show="headings", selectmode="browse")
        self.tree.grid(row=1, column=0, sticky="nsew", padx=10, pady=(0, 10))
        setup = [
            ("code", "Codigo", 120),
            ("name", "Nombre", 260),
            ("brand", "Marca", 120),
            ("category", "Categoria", 130),
            ("stock", "Inventario", 90),
            ("cost", "Costo", 90),
            ("public", "$PUBLICO", 90),
            ("tax", "IVA", 70),
        ]
        for column, title, width in setup:
            self.tree.heading(column, text=title)
            self.tree.column(column, width=width, anchor="e" if column in {"stock", "cost", "public", "tax"} else "w")
        self.tree.bind("<Double-1>", lambda _event: self.edit_product())

        bottom = ttk.Frame(self, padding=(10, 0, 10, 10))
        bottom.grid(row=2, column=0, sticky="ew")
        for col in range(6):
            bottom.columnconfigure(col, weight=1)
        ttk.Button(bottom, text="Nuevo", command=self.new_product).grid(row=0, column=0, sticky="ew", padx=4)
        ttk.Button(bottom, text="Editar", command=self.edit_product).grid(row=0, column=1, sticky="ew", padx=4)
        ttk.Button(bottom, text="Eliminar", command=self.delete_product).grid(row=0, column=2, sticky="ew", padx=4)
        ttk.Button(bottom, text="Actualizar", command=self.search).grid(row=0, column=3, sticky="ew", padx=4)
        ttk.Button(bottom, text="Excel", command=self.export_inventory).grid(row=0, column=4, sticky="ew", padx=4)
        ttk.Button(bottom, text="Cerrar", command=self.destroy).grid(row=0, column=5, sticky="ew", padx=4)

    def search(self) -> None:
        for row in self.tree.get_children():
            self.tree.delete(row)
        rows = search_products(self.conn, self.query_var.get(), limit=500)
        for row in rows:
            self.tree.insert(
                "",
                "end",
                iid=row["code"],
                values=(
                    row["code"],
                    row["name"],
                    row["brand"] or "",
                    row["category"] or "",
                    number(row["stock"] or 0),
                    money(row["cost"]),
                    money(row["price_public"]),
                    f"{row['tax_rate'] or 0:g}%",
                ),
            )

    def import_inventory(self) -> None:
        if self.on_import is not None:
            self.on_import()
            self.search()

    def export_inventory(self) -> None:
        path = filedialog.asksaveasfilename(
            parent=self,
            title="Exportar inventario a Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile="Inventario_AREZONE.xlsx",
        )
        if not path:
            return

        rows = self.conn.execute(
            """
            SELECT code, name, brand_code, brand, group_code, category, alt_codes, stock, cost, price_public, tax_rate, description
            FROM products
            ORDER BY name
            """
        ).fetchall()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Inventario"
        headers = [
            "Codigo",
            "Nombre",
            "Cod Marca",
            "Marca",
            "Cod Grupo",
            "Grupo",
            "Codigos Alternos",
            "Inventario",
            "Costo",
            "Precio Publico",
            "IVA %",
            "Descripcion",
        ]
        sheet.append(headers)

        for row in rows:
            sheet.append(
                [
                    row["code"],
                    row["name"],
                    row["brand_code"] or "",
                    row["brand"] or "",
                    row["group_code"] or "",
                    row["category"] or "",
                    row["alt_codes"] or "",
                    float(row["stock"] or 0),
                    float(row["cost"] or 0),
                    float(row["price_public"] or 0),
                    float(row["tax_rate"] or 0),
                    row["description"] or "",
                ]
            )

        header_fill = PatternFill("solid", fgColor="1F2937")
        header_font = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="D1D5DB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        widths = {
            "A": 18,
            "B": 34,
            "C": 14,
            "D": 20,
            "E": 14,
            "F": 20,
            "G": 28,
            "H": 12,
            "I": 14,
            "J": 16,
            "K": 10,
            "L": 36,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        workbook.save(path)
        messagebox.showinfo("Exportar", f"Inventario exportado en:\n{path}")

    def selected_code(self) -> str | None:
        selected = self.tree.selection()
        return selected[0] if selected else None

    def new_product(self) -> None:
        dialog = ProductFormDialog(self, self.conn)
        self.wait_window(dialog)
        if dialog.saved:
            self.search()

    def edit_product(self) -> None:
        code = self.selected_code()
        if not code:
            messagebox.showinfo("Inventario", "Seleccione un producto.")
            return
        row = self.conn.execute("SELECT * FROM products WHERE code = ?", (code,)).fetchone()
        if not row:
            return
        dialog = ProductFormDialog(self, self.conn, dict(row))
        self.wait_window(dialog)
        if dialog.saved:
            self.search()

    def delete_product(self) -> None:
        code = self.selected_code()
        if not code:
            messagebox.showinfo("Inventario", "Seleccione un producto.")
            return
        master = ask_secret(self, "Clave maestra", "Ingrese clave maestra para eliminar:", show="*")
        if master is None:
            return
        master_hash = get_setting(self.conn, "master_hash")
        if not master_hash or not verify_secret(master, master_hash):
            messagebox.showerror("Clave maestra", "Clave maestra incorrecta.")
            return
        if not messagebox.askyesno("Eliminar", f"Eliminar producto {code}?"):
            return
        self.conn.execute("DELETE FROM products WHERE code = ?", (code,))
        self.conn.commit()
        self.search()


class ProductFormDialog(tk.Toplevel):
    def __init__(self, master, conn: sqlite3.Connection, product: dict | None = None) -> None:
        super().__init__(master)
        self.conn = conn
        self.product = product
        self.saved = False
        self.original_code = (product or {}).get("code", "")

        self.vars = {
            "code": tk.StringVar(value=(product or {}).get("code", "")),
            "alt_codes": tk.StringVar(value=(product or {}).get("alt_codes", "")),
            "name": tk.StringVar(value=(product or {}).get("name", "")),
            "brand_code": tk.StringVar(value=(product or {}).get("brand_code", "")),
            "brand": tk.StringVar(value=(product or {}).get("brand", "")),
            "group_code": tk.StringVar(value=(product or {}).get("group_code", "")),
            "category": tk.StringVar(value=(product or {}).get("category", "")),
            "description": tk.StringVar(value=(product or {}).get("description", "")),
            "stock": tk.StringVar(value=str((product or {}).get("stock", 0))),
            "cost": tk.StringVar(value=str((product or {}).get("cost", 0))),
            "price_public": tk.StringVar(value=str((product or {}).get("price_public", 0))),
            "tax_rate": tk.StringVar(value=str((product or {}).get("tax_rate", 0))),
        }
        self.has_tax_var = tk.IntVar(value=int((product or {}).get("has_tax", 0) or 0))

        self.title("Producto")
        fit_window(self, 640, 680, min_width=580, min_height=620)
        self.transient(master)
        self.grab_set()
        self.columnconfigure(0, weight=1)

        self._build()
        self._load_catalogs()

    def _build(self) -> None:
        form = ttk.Frame(self, padding=16)
        form.grid(row=0, column=0, sticky="nsew")
        form.columnconfigure(1, weight=1)
        form.columnconfigure(2, weight=1)

        fields = [
            ("Codigo", "code"),
            ("Codigos alternos", "alt_codes"),
            ("Nombre", "name"),
            ("Cod marca", "brand_code"),
            ("Marca", "brand"),
            ("Cod grupo", "group_code"),
            ("Categoria", "category"),
            ("Descripcion", "description"),
            ("Inventario", "stock"),
            ("Costo", "cost"),
            ("$PUBLICO", "price_public"),
            ("IVA %", "tax_rate"),
        ]
        for row, (label, key) in enumerate(fields):
            ttk.Label(form, text=label).grid(row=row, column=0, sticky="w", pady=4)
            if key in {"brand", "category"}:
                combo = ttk.Combobox(form, textvariable=self.vars[key], state="normal")
                combo.grid(row=row, column=1, sticky="ew", pady=4)
                if key == "brand":
                    self.brand_combo = combo
                else:
                    self.category_combo = combo
            else:
                ttk.Entry(form, textvariable=self.vars[key]).grid(
                    row=row,
                    column=1,
                    sticky="ew",
                    pady=4,
                )
            if key == "alt_codes":
                tk.Label(form, text="Primero queda como principal. Separa con coma.", fg="#666666").grid(
                    row=row,
                    column=2,
                    sticky="w",
                    padx=(8, 0),
                )
        ttk.Checkbutton(form, text="Con IVA", variable=self.has_tax_var).grid(
            row=len(fields),
            column=1,
            sticky="w",
            pady=8,
        )

        actions = ttk.Frame(form)
        actions.grid(row=len(fields) + 1, column=0, columnspan=2, sticky="ew", pady=(16, 0))
        actions.columnconfigure((0, 1), weight=1)
        ttk.Button(actions, text="Guardar", command=self.save).grid(row=0, column=0, sticky="ew", padx=(0, 5))
        ttk.Button(actions, text="Cancelar", command=self.destroy).grid(row=0, column=1, sticky="ew", padx=(5, 0))

    def _load_catalogs(self) -> None:
        brands = [row["name"] for row in self.conn.execute("SELECT name FROM product_brands ORDER BY name").fetchall()]
        categories = [row["name"] for row in self.conn.execute("SELECT name FROM product_categories ORDER BY name").fetchall()]
        if hasattr(self, "brand_combo"):
            self.brand_combo["values"] = brands
        if hasattr(self, "category_combo"):
            self.category_combo["values"] = categories

    def save(self) -> None:
        code, alt_codes = split_primary_and_alternates(self.vars["code"].get(), self.vars["alt_codes"].get())
        name = self.vars["name"].get().strip()
        if not code or not name:
            messagebox.showwarning("Producto", "Codigo y nombre son obligatorios.")
            return

        now = now_text()
        values = {key: var.get().strip() for key, var in self.vars.items()}
        values["code"] = code
        values["name"] = name
        numeric = {
            "cost": parse_amount(values["cost"]),
            "stock": parse_amount(values["stock"]),
            "price_public": parse_amount(values["price_public"]),
            "tax_rate": parse_amount(values["tax_rate"]),
        }

        if values["brand"]:
            self.conn.execute("INSERT OR IGNORE INTO product_brands(name) VALUES(?)", (values["brand"],))
        if values["category"]:
            self.conn.execute("INSERT OR IGNORE INTO product_categories(name) VALUES(?)", (values["category"],))

        if self.product:
            if code != self.original_code:
                exists = self.conn.execute("SELECT code FROM products WHERE code = ?", (code,)).fetchone()
                if exists:
                    messagebox.showwarning("Producto", "La nueva referencia ya existe.")
                    return
            self.conn.execute(
                """
                UPDATE products
                SET code = ?, alt_codes = ?, name = ?, brand_code = ?, brand = ?, group_code = ?, category = ?,
                    stock = ?, has_tax = ?, tax_rate = ?, description = ?, cost = ?,
                    price_public = ?, price_1 = ?, price_2 = ?, updated_at = ?
                WHERE code = ?
                """,
                (
                    code,
                    alt_codes,
                    name,
                    values["brand_code"],
                    values["brand"],
                    values["group_code"],
                    values["category"],
                    numeric["stock"],
                    self.has_tax_var.get(),
                    numeric["tax_rate"],
                    values["description"],
                    numeric["cost"],
                    numeric["price_public"],
                    numeric["price_public"],
                    numeric["price_public"],
                    now,
                    self.original_code,
                ),
            )
        else:
            existing = self.conn.execute("SELECT code FROM products WHERE code = ?", (code,)).fetchone()
            if existing:
                messagebox.showwarning("Producto", "Ya existe un producto con ese codigo.")
                return
            self.conn.execute(
                """
                INSERT INTO products (
                    code, alt_codes, name, brand_code, brand, group_code, category, stock, has_tax,
                    tax_rate, description, cost, price_public, price_1, price_2,
                    created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    code,
                    alt_codes,
                    name,
                    values["brand_code"],
                    values["brand"],
                    values["group_code"],
                    values["category"],
                    numeric["stock"],
                    self.has_tax_var.get(),
                    numeric["tax_rate"],
                    values["description"],
                    numeric["cost"],
                    numeric["price_public"],
                    numeric["price_public"],
                    numeric["price_public"],
                    now,
                    now,
                ),
            )
        self.conn.commit()
        self.saved = True
        self.destroy()
