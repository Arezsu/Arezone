from __future__ import annotations

import re
import sqlite3
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any


@dataclass
class MigrationResult:
    imported: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[str] | None = None

    @property
    def total_ok(self) -> int:
        return self.imported + self.updated


FIELD_MAP = {
    "REFERENCIA": "code",
    "DETALLE": "name",
    "COD MARCA": "brand_code",
    "MARCA": "brand",
    "COD GRUPO": "group_code",
    "GRUPO": "category",
    "CON IVA": "has_tax",
    "IVA": "tax_rate",
    "DESCRIPCION": "description",
    "EXISTENCIAS": "stock",
    "STOCK": "stock",
    "INVENTARIO": "stock",
    "COSTO": "cost",
    "$PUBLICO": "price_public",
    "%PUBLICO": "margin_public",
    "$PRECIO1": "price_1",
    "%PRECIO1": "margin_1",
    "$PRECIO2": "price_2",
    "%PRECIO2": "margin_2",
}

NUMERIC_FIELDS = {"tax_rate", "stock", "cost", "price_public", "price_1", "price_2"}


def migrate_inventory(excel_path: str | Path, conn: sqlite3.Connection) -> MigrationResult:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - depends on user environment
        return MigrationResult(errors=[f"No se pudo cargar openpyxl: {exc}"])

    path = Path(excel_path)
    if not path.exists():
        return MigrationResult(errors=[f"No existe el archivo: {path}"])

    result = MigrationResult(errors=[])
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        return MigrationResult(errors=[f"No se pudo abrir el Excel: {exc}"])

    sheet = workbook["RPTARJETA_TO_XLS"] if "RPTARJETA_TO_XLS" in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration:
        return MigrationResult(errors=["El archivo no tiene filas."])

    header_map = _build_header_map(raw_headers)
    if "REFERENCIA" not in header_map or "DETALLE" not in header_map:
        return MigrationResult(
            errors=["El Excel debe tener al menos las columnas REFERENCIA y DETALLE."]
        )

    now = datetime.now().isoformat(timespec="seconds")
    for row_index, raw_row in enumerate(rows, start=2):
        try:
            data = _row_to_product(raw_row, header_map)
            if not data["code"] or not data["name"]:
                result.skipped += 1
                continue

            existing = conn.execute(
                "SELECT code FROM products WHERE code = ?",
                (data["code"],),
            ).fetchone()

            values = (
                data["code"],
                data["name"],
                data.get("brand_code"),
                data.get("brand"),
                data.get("group_code"),
                data.get("category"),
                data.get("stock", 0),
                data.get("has_tax", 0),
                data.get("tax_rate", 0),
                data.get("description"),
                data.get("cost", 0),
                data.get("price_public", 0),
                data.get("price_1", 0),
                data.get("price_2", 0),
                now,
                now,
            )
            conn.execute(
                """
                INSERT INTO products (
                    code, name, brand_code, brand, group_code, category, stock, has_tax,
                    tax_rate, description, cost, price_public, price_1, price_2,
                    created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(code) DO UPDATE SET
                    name = excluded.name,
                    brand_code = excluded.brand_code,
                    brand = excluded.brand,
                    group_code = excluded.group_code,
                    category = excluded.category,
                    stock = excluded.stock,
                    has_tax = excluded.has_tax,
                    tax_rate = excluded.tax_rate,
                    description = excluded.description,
                    cost = excluded.cost,
                    price_public = excluded.price_public,
                    price_1 = excluded.price_1,
                    price_2 = excluded.price_2,
                    updated_at = excluded.updated_at
                """,
                values,
            )
            if existing:
                result.updated += 1
            else:
                result.imported += 1
        except Exception as exc:
            result.skipped += 1
            if result.errors is not None and len(result.errors) < 8:
                result.errors.append(f"Fila {row_index}: {exc}")

    conn.commit()
    return result


def _build_header_map(raw_headers: tuple[Any, ...]) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for index, value in enumerate(raw_headers):
        normalized = _normalize_header(value)
        if normalized:
            header_map[normalized] = index
    return header_map


def _row_to_product(row: tuple[Any, ...], header_map: dict[str, int]) -> dict[str, Any]:
    data: dict[str, Any] = {
        "code": "",
        "name": "",
        "brand_code": "",
        "brand": "",
        "group_code": "",
        "category": "",
        "stock": 0,
        "has_tax": 0,
        "tax_rate": 0,
        "description": "",
        "cost": 0,
        "price_public": 0,
        "price_1": 0,
        "price_2": 0,
    }

    for source_name, target_name in FIELD_MAP.items():
        if source_name not in header_map or target_name.startswith("margin"):
            continue
        value = row[header_map[source_name]] if header_map[source_name] < len(row) else None
        if target_name == "has_tax":
            data[target_name] = _to_bool(value)
        elif target_name in NUMERIC_FIELDS:
            data[target_name] = _to_money(value)
        else:
            data[target_name] = _clean_text(value)

    data["code"] = _clean_code(data["code"])
    if not data["price_1"]:
        data["price_1"] = data["price_public"]
    if not data["price_2"]:
        data["price_2"] = data["price_public"]
    return data


def _normalize_header(value: Any) -> str:
    text = _clean_text(value).upper()
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _clean_code(value: Any) -> str:
    code = _clean_text(value)
    if code.endswith(".0"):
        code = code[:-2]
    return code


def _to_bool(value: Any) -> int:
    text = _clean_text(value).upper()
    return 1 if text in {"SI", "S", "YES", "Y", "1", "TRUE", "VERDADERO"} else 0


def _to_money(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    text = _clean_text(value)
    text = text.replace("$", "").replace("COP", "").replace(" ", "")
    text = re.sub(r"[^0-9,.\-]", "", text)
    if not text:
        return 0.0

    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(".", "").replace(",", ".")

    try:
        return float(text)
    except ValueError:
        return 0.0
